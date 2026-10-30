"""Sinh file Excel mẫu theo cơ cấu phòng thật của trường: docs/samples/school_realistic.xlsx.

Chạy:  uv run python scripts/make_school_sample.py

Dữ liệu cố định (không ngẫu nhiên) nên chạy lại luôn ra cùng một file.
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path("docs/samples/school_realistic.xlsx")
OUT_SMALL = Path("docs/samples/school_demo_small.xlsx")
# Bản nhỏ để thử nhanh trên máy yếu: mỗi hướng một lớp, kèm lớp MBA và lớp Design.
SMALL_CLASSES = {"IT1", "BA1", "MK1", "AC1", "MBA1", "DS1"}
CAPACITY = 30  # sức chứa tạm thời, chờ số chính xác (xem docs/requirements.md)

# (id, tên, loại chính, hỗ trợ thêm). Chỉ gồm phòng có dùng để xếp lớp.
ROOMS = [
    ("G313", "Phòng hội họa", "art", ""),
    ("G314", "Phòng G314", "standard", ""),
    ("G316", "Phòng G316 (MBA)", "mba", "standard"),
    ("G317", "Phòng G317", "standard", ""),
    ("G318", "Phòng G318", "standard", ""),
    ("G319", "Phòng G319", "standard", ""),
    ("G502", "Phòng G502", "standard", ""),
    ("G503", "Phòng G503", "standard", ""),
    ("G504", "Phòng G504", "standard", ""),
    ("G508", "Phòng G508", "standard", ""),
    ("G509", "Phòng G509", "standard", ""),
    ("G510", "Phòng G510", "standard", ""),
]

# Phòng không đưa vào bài toán và lý do (ghi ở sheet Ghi chú).
EXCLUDED = [
    ("G311", "Hội trường, tổ chức sự kiện"),
    ("G312", "Hội trường, tổ chức sự kiện"),
    ("G315", "Phòng tư vấn tuyển sinh"),
    ("G320", "Phòng để đồ, chưa có thông tin"),
    ("G501", "Phòng Director"),
    ("G505", "Academic Affairs Administration Dept."),
    ("G506", "Thư viện"),
    ("G507", "Student service"),
]

TEACHERS = [f"GV{i}" for i in range(1, 21)]

# Mỗi hướng học có 6 môn thường: (tên môn, số tiết/tuần).
TRACKS = {
    "IT": [
        ("Lập trình", 4),
        ("Cơ sở dữ liệu", 3),
        ("Mạng máy tính", 3),
        ("Toán rời rạc", 3),
        ("Tiếng Anh", 3),
        ("Kỹ năng mềm", 2),
    ],
    "BA": [
        ("Quản trị học", 3),
        ("Kinh tế vi mô", 3),
        ("Marketing căn bản", 3),
        ("Kế toán", 3),
        ("Tiếng Anh", 4),
        ("Kỹ năng mềm", 2),
    ],
    "MK": [
        ("Hành vi người tiêu dùng", 3),
        ("Truyền thông số", 4),
        ("Nghiên cứu thị trường", 3),
        ("Quản trị thương hiệu", 3),
        ("Tiếng Anh", 3),
        ("Kỹ năng mềm", 2),
    ],
    "AC": [
        ("Nguyên lý kế toán", 4),
        ("Kế toán tài chính", 3),
        ("Thuế", 3),
        ("Kiểm toán căn bản", 3),
        ("Tiếng Anh", 3),
        ("Kỹ năng mềm", 2),
    ],
}
MBA_COURSES = [
    ("Lãnh đạo", 4),
    ("Chiến lược kinh doanh", 4),
    ("Tài chính doanh nghiệp", 3),
    ("Quản trị vận hành", 3),
]
DESIGN_COURSES = [
    ("Vẽ cơ bản", 6, "art"),
    ("Thiết kế đồ họa", 3, "standard"),
    ("Lịch sử mỹ thuật", 3, "standard"),
    ("Tiếng Anh", 4, "standard"),
    ("Kỹ năng mềm", 2, "standard"),
]


def build(only: set[str] | None = None) -> tuple[list[tuple], list[tuple]]:
    classes: list[tuple] = []
    courses: list[tuple] = []  # (class_id, tên, số tiết, loại phòng)
    for track, plan in TRACKS.items():
        for n, size in ((1, 30), (2, 28)):
            cid = f"{track}{n}"
            classes.append((cid, f"Lớp {cid}", size, "regular"))
            courses += [(cid, name, periods, "standard") for name, periods in plan]
    for n, size in ((1, 24), (2, 20)):
        cid = f"MBA{n}"
        classes.append((cid, f"Lớp {cid}", size, "mba"))
        courses += [(cid, name, periods, "mba") for name, periods in MBA_COURSES]
    for n, size in ((1, 26), (2, 24)):
        cid = f"DS{n}"
        classes.append((cid, f"Lớp {cid}", size, "design"))
        courses += [(cid, name, periods, room) for name, periods, room in DESIGN_COURSES]

    if only is not None:
        classes = [c for c in classes if c[0] in only]
        courses = [c for c in courses if c[0] in only]

    # Chia giáo viên: môn nhiều tiết trước, mỗi môn cho giáo viên đang ít tiết nhất
    # (và chưa dạy lớp đó), để tải đều giữa 20 giáo viên.
    load = {t: 0 for t in TEACHERS}
    taught: dict[str, set[str]] = {t: set() for t in TEACHERS}
    assigned: dict[int, str] = {}
    order = sorted(range(len(courses)), key=lambda i: (-courses[i][2], i))
    for i in order:
        cid = courses[i][0]
        teacher = min((t for t in TEACHERS if cid not in taught[t]), key=lambda t: (load[t], t))
        assigned[i] = teacher
        load[teacher] += courses[i][2]
        taught[teacher].add(cid)
    rows = [
        (f"C{i + 1:03d}", name, cid, assigned[i], periods, room)
        for i, (cid, name, periods, room) in enumerate(courses)
    ]
    return classes, rows


def write(out: Path, only: set[str] | None = None) -> None:
    classes, courses = build(only)
    wb = Workbook()
    ws = wb.active
    ws.title = "Phòng"
    ws.append(["id", "name", "capacity", "room_type", "also_supports"])
    for room_id, name, room_type, also in ROOMS:
        ws.append([room_id, name, CAPACITY, room_type, also or None])
    ws = wb.create_sheet("Giáo viên")
    ws.append(["id", "name"])
    for t in TEACHERS:
        ws.append([t, t])
    ws = wb.create_sheet("Lớp")
    ws.append(["id", "name", "size", "kind"])
    for row in classes:
        ws.append(list(row))
    ws = wb.create_sheet("Môn")
    ws.append(["id", "name", "class_id", "teacher_id", "periods_per_week", "required_room_type"])
    for row in courses:
        ws.append(list(row))
    ws = wb.create_sheet("Ghi chú")
    ws.append(["Mục", "Nội dung"])
    ws.append(
        ["Mục đích", "Dữ liệu mẫu theo cơ cấu phòng thật. Sinh bằng scripts/make_school_sample.py."]
    )
    ws.append(["Lưới giờ", "Thứ 2 đến Thứ 7, tiết 1 đến 6 (tiết 7 và 8 không xếp)."])
    ws.append(
        [
            "G316",
            "Phòng học thường; lớp MBA bắt buộc dùng phòng này (loại mba, hỗ trợ thêm standard).",
        ]
    )
    ws.append(["G313", "Phòng hội họa, chỉ môn vẽ (loại art)."])
    ws.append(["Sức chứa", f"Tạm đặt {CAPACITY} cho mọi phòng, chờ số chính xác."])
    for room_id, why in EXCLUDED:
        ws.append([f"{room_id} (loại khỏi bài toán)", why])
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(
        f"Đã ghi {out}: {len(ROOMS)} phòng, {len(TEACHERS)} giáo viên, "
        f"{len(classes)} lớp, {len(courses)} môn, {sum(c[4] for c in courses)} tiết/tuần"
    )


def main() -> None:
    write(OUT)
    write(OUT_SMALL, SMALL_CLASSES)


if __name__ == "__main__":
    main()
