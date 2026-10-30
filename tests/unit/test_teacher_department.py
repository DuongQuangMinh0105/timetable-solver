from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook

from timetable.domain.services.teacher_department import check_teacher_departments
from timetable.infrastructure.excel.importer import OpenpyxlExcelImporter
from timetable.infrastructure.problem_loader import InvalidProblemError, problem_from_dict
from timetable.interfaces.ui.tables import payload_to_tables, tables_to_payload


def _data(course_dept: str | None, teacher_dept: str | None) -> dict[str, Any]:
    return {
        "schema_version": "1.0",
        "grid": {"days": 6, "periods": 6, "allowed_periods": [0, 1, 2, 3, 4, 5]},
        "rooms": [{"id": "R1", "name": "R1", "capacity": 30}],
        "teachers": [
            {"id": "GV1", "name": "GV1", "department": teacher_dept},
            {"id": "GV2", "name": "GV2", "department": "IT"},
        ],
        "class_groups": [{"id": "C1", "name": "C1", "size": 20}],
        "courses": [
            {
                "id": "M1",
                "name": "Lập trình",
                "class_id": "C1",
                "teacher_id": "GV1",
                "periods_per_week": 2,
                "department": course_dept,
            }
        ],
    }


def test_matching_department_is_accepted() -> None:
    assert check_teacher_departments(problem_from_dict(_data("IT", "IT"))) == []


def test_comparison_ignores_case_and_spaces() -> None:
    assert check_teacher_departments(problem_from_dict(_data("it", " IT "))) == []


@pytest.mark.parametrize(("course", "teacher"), [(None, "Design"), ("IT", None), (None, None)])
def test_missing_department_means_no_restriction(course: str | None, teacher: str | None) -> None:
    assert check_teacher_departments(problem_from_dict(_data(course, teacher))) == []


def test_mismatch_is_reported_in_vietnamese_with_suggestion() -> None:
    with pytest.raises(InvalidProblemError) as error:
        problem_from_dict(_data("IT", "Design"))
    text = str(error.value)
    assert "teacher_department_mismatch" in text
    assert "Lập trình" in text and "GV1" in text and "Design" in text
    assert "GV2" in text  # gợi ý giáo viên nhóm IT


def _workbook(path: Path, course_dept: str, teacher_dept: str) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Phòng"
    ws.append(["id", "name", "capacity"])
    ws.append(["R1", "R1", 30])
    wb.create_sheet("Giáo viên").append(["id", "name", "department"])
    wb["Giáo viên"].append(["GV1", "GV1", teacher_dept])
    wb.create_sheet("Lớp").append(["id", "name", "size"])
    wb["Lớp"].append(["C1", "C1", 20])
    wb.create_sheet("Môn").append(
        ["id", "name", "class_id", "teacher_id", "periods_per_week", "department"]
    )
    wb["Môn"].append(["M1", "Lập trình", "C1", "GV1", 2, course_dept])
    wb.save(path)
    return path


def test_excel_importer_reads_department_columns(tmp_path: Path) -> None:
    data = OpenpyxlExcelImporter().load_data(_workbook(tmp_path / "a.xlsx", "IT", "IT"))
    assert data["teachers"][0]["department"] == "IT"
    assert data["courses"][0]["department"] == "IT"
    assert problem_from_dict(data).courses[0].department == "IT"


def test_excel_without_department_columns_still_loads(tmp_path: Path) -> None:
    path = _workbook(tmp_path / "b.xlsx", "", "")
    data = OpenpyxlExcelImporter().load_data(path)
    assert data["teachers"][0]["department"] is None
    assert data["courses"][0]["department"] is None


def test_ui_tables_keep_department_and_report_mismatch() -> None:
    payload = _data("IT", "IT")
    tables = payload_to_tables(payload)
    assert tables["teachers"][0]["department"] == "IT"
    rebuilt, errors = tables_to_payload(tables, payload)
    assert errors == [] and rebuilt is not None
    assert rebuilt["courses"][0]["department"] == "IT"

    tables["teachers"][0]["department"] = "Design"
    rebuilt, errors = tables_to_payload(tables, payload)
    assert rebuilt is None and any("teacher_department_mismatch" in e for e in errors)
