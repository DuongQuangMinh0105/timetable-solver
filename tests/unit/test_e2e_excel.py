from pathlib import Path

from excel_helpers import sheets_from_data, write_workbook
from openpyxl import load_workbook

from timetable.application.dto.solve import SolveOptions, SolveStatus
from timetable.domain.services.validation_report import build_report
from timetable.infrastructure.excel.exporter import OpenpyxlExcelExporter
from timetable.infrastructure.excel.importer import OpenpyxlExcelImporter
from timetable.infrastructure.generator import generate_v1
from timetable.infrastructure.solver.default_solver import build_default_solver


def test_import_solve_export_validate(tmp_path: Path) -> None:
    data = generate_v1()
    source = write_workbook(tmp_path / "in.xlsx", sheets_from_data(data))

    problem = OpenpyxlExcelImporter().load(source)
    solver = build_default_solver(problem, data)
    result = solver.solve(problem, SolveOptions(time_limit_s=30, num_workers=2))

    assert result.status in (SolveStatus.OPTIMAL, SolveStatus.FEASIBLE)
    assert result.timetable is not None

    out = tmp_path / "out.xlsx"
    OpenpyxlExcelExporter().export(problem, result.timetable, out)
    assert load_workbook(out).sheetnames == ["Theo lớp", "Theo giáo viên", "Theo phòng"]

    report = build_report(problem, result.timetable)
    assert report.hard == ()
    assert report.summary()["hard"] == 0
