from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from timetable.application.ports.excel_exporter import ExcelExportError
from timetable.domain.entities import Problem
from timetable.domain.entities.timetable import Assignment, Timetable
from timetable.infrastructure.excel.exporter import OpenpyxlExcelExporter
from timetable.infrastructure.generator import generate_v1
from timetable.infrastructure.problem_loader import problem_from_dict


def _problem() -> Problem:
    return problem_from_dict(generate_v1())


def _first_cell_row(problem: Problem, index: int) -> int:
    """Dòng Excel của tiết đầu tiên trong khối thứ `index`."""
    block = len(problem.grid.allowed_periods) + 3
    return index * block + 3


def _text(ws: Any, row: int, day: int) -> str:
    return str(ws.cell(row=row, column=day + 2).value or "")


def test_export_creates_the_three_sheets(tmp_path: Path) -> None:
    problem = _problem()
    out = tmp_path / "out.xlsx"
    OpenpyxlExcelExporter().export(problem, Timetable(), out)
    assert load_workbook(out).sheetnames == ["Theo lớp", "Theo giáo viên", "Theo phòng"]


def test_class_view_shows_subject_teacher_and_room(tmp_path: Path) -> None:
    problem = _problem()
    unit = problem.lesson_units[0]
    room = problem.rooms[0]
    out = tmp_path / "out.xlsx"
    timetable = Timetable((Assignment(unit.id, 1, 0, room.id),))
    OpenpyxlExcelExporter().export(problem, timetable, out)

    course = next(c for c in problem.courses if c.id == unit.course_id)
    teacher = next(t for t in problem.teachers if t.id == unit.teacher_id)
    index = [c.id for c in problem.class_groups].index(unit.class_id)
    cell = _text(load_workbook(out)["Theo lớp"], _first_cell_row(problem, index), 1)
    assert course.name in cell
    assert teacher.name in cell
    assert room.name in cell


def test_teacher_and_room_views_show_the_lesson(tmp_path: Path) -> None:
    problem = _problem()
    unit = problem.lesson_units[0]
    room_index = 2
    room = problem.rooms[room_index]
    out = tmp_path / "out.xlsx"
    OpenpyxlExcelExporter().export(problem, Timetable((Assignment(unit.id, 0, 0, room.id),)), out)
    wb = load_workbook(out)

    class_name = next(c.name for c in problem.class_groups if c.id == unit.class_id)
    teacher_index = [t.id for t in problem.teachers].index(unit.teacher_id)
    assert class_name in _text(wb["Theo giáo viên"], _first_cell_row(problem, teacher_index), 0)
    assert class_name in _text(wb["Theo phòng"], _first_cell_row(problem, room_index), 0)


def test_two_period_unit_fills_both_periods(tmp_path: Path) -> None:
    data = generate_v1()
    data["lesson_units"] = []  # để loader tự tách môn thành block 2 + 1 tiết
    problem = problem_from_dict(data)
    unit = next(u for u in problem.lesson_units if u.duration == 2)
    out = tmp_path / "out.xlsx"
    OpenpyxlExcelExporter().export(
        problem, Timetable((Assignment(unit.id, 0, 0, problem.rooms[0].id),)), out
    )
    index = [c.id for c in problem.class_groups].index(unit.class_id)
    ws = load_workbook(out)["Theo lớp"]
    first = _first_cell_row(problem, index)
    assert _text(ws, first, 0) != ""
    assert _text(ws, first + 1, 0) == _text(ws, first, 0)


def test_unknown_unit_is_reported(tmp_path: Path) -> None:
    problem = _problem()
    timetable = Timetable((Assignment("KHONG-CO", 0, 0, problem.rooms[0].id),))
    with pytest.raises(ExcelExportError, match="KHONG-CO"):
        OpenpyxlExcelExporter().export(problem, timetable, tmp_path / "out.xlsx")
