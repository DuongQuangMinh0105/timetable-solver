from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook  # type: ignore[import-untyped]
from openpyxl.styles import Alignment, Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

from timetable.application.ports.excel_exporter import ExcelExportError
from timetable.domain.entities import Problem, Timetable

SHEET_BY_CLASS = "Theo lớp"
SHEET_BY_TEACHER = "Theo giáo viên"
SHEET_BY_ROOM = "Theo phòng"

DAY_NAMES = ("Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "Chủ nhật")

# owner_id -> (day, period) -> các dòng chữ trong ô
_Cells = dict[str, dict[tuple[int, int], list[str]]]


def _day_name(day: int) -> str:
    return DAY_NAMES[day] if day < len(DAY_NAMES) else f"Ngày {day + 1}"


def _write_sheet(
    wb: Any, title: str, owners: list[tuple[str, str]], cells: _Cells, problem: Problem
) -> None:
    ws = wb.create_sheet(title)
    wrap = Alignment(wrap_text=True, vertical="top")
    header_fill = PatternFill("solid", fgColor="D9E1F2")
    days = problem.grid.days
    row = 1
    for owner_id, owner_name in owners:
        ws.cell(row=row, column=1, value=f"{owner_name} ({owner_id})").font = Font(
            bold=True, size=12
        )
        row += 1
        for col, text in enumerate(["Tiết", *(_day_name(d) for d in range(days))], start=1):
            cell = ws.cell(row=row, column=col, value=text)
            cell.font = Font(bold=True)
            cell.fill = header_fill
        row += 1
        for period in problem.grid.allowed_periods:
            ws.cell(row=row, column=1, value=period + 1)
            for day in range(days):
                lines = cells.get(owner_id, {}).get((day, period), [])
                ws.cell(row=row, column=day + 2, value="\n".join(lines) or None).alignment = wrap
            row += 1
        row += 1  # dòng trống giữa các khối
    ws.column_dimensions["A"].width = 8
    for day in range(days):
        ws.column_dimensions[get_column_letter(day + 2)].width = 26
    ws.freeze_panes = "B1"


class OpenpyxlExcelExporter:
    def export(self, problem: Problem, timetable: Timetable, path: Path) -> None:
        units = {u.id: u for u in problem.lesson_units}
        courses = {c.id: c for c in problem.courses}
        rooms = {r.id: r for r in problem.rooms}
        teachers = {t.id: t for t in problem.teachers}
        classes = {c.id: c for c in problem.class_groups}

        by_class: _Cells = defaultdict(lambda: defaultdict(list))
        by_teacher: _Cells = defaultdict(lambda: defaultdict(list))
        by_room: _Cells = defaultdict(lambda: defaultdict(list))

        for a in timetable.assignments:
            unit = units.get(a.unit_id)
            if unit is None:
                raise ExcelExportError(f"Lịch tham chiếu đơn vị học không có trong bài toán: {a}")
            course = courses.get(unit.course_id)
            subject = course.name if course else unit.course_id
            room = rooms.get(a.room_id)
            teacher = teachers.get(unit.teacher_id)
            class_group = classes.get(unit.class_id)
            room_name = room.name if room else a.room_id
            teacher_name = teacher.name if teacher else unit.teacher_id
            class_name = class_group.name if class_group else unit.class_id

            for period in range(a.period, a.period + unit.duration):
                key = (a.day, period)
                by_class[unit.class_id][key].append(f"{subject}\n{teacher_name} - {room_name}")
                by_teacher[unit.teacher_id][key].append(f"{subject}\n{class_name} - {room_name}")
                by_room[a.room_id][key].append(f"{subject}\n{class_name} - {teacher_name}")

        wb = Workbook()
        wb.remove(wb.active)
        _write_sheet(
            wb, SHEET_BY_CLASS, [(c.id, c.name) for c in problem.class_groups], by_class, problem
        )
        _write_sheet(
            wb, SHEET_BY_TEACHER, [(t.id, t.name) for t in problem.teachers], by_teacher, problem
        )
        _write_sheet(wb, SHEET_BY_ROOM, [(r.id, r.name) for r in problem.rooms], by_room, problem)
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
