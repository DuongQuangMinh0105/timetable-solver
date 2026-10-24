import zipfile
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from timetable.application.ports.excel_importer import ExcelImportError
from timetable.domain.entities import Problem
from timetable.infrastructure.excel.english_schema import convert, is_english_workbook
from timetable.infrastructure.generator import (
    ALLOWED_PERIODS,
    MAX_CONSECUTIVE_PERIODS,
    NUM_DAYS,
    NUM_PERIODS,
    SCHEMA_VERSION,
)
from timetable.infrastructure.problem_loader import problem_from_dict

SHEET_ROOMS = "Phòng"
SHEET_TEACHERS = "Giáo viên"
SHEET_CLASSES = "Lớp"
SHEET_COURSES = "Môn"

_Row = tuple[int, dict[str, Any]]


def _clean(value: Any) -> Any:
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _read_sheet(
    wb: Any, sheet: str, required: tuple[str, ...], optional: tuple[str, ...], errors: list[str]
) -> list[_Row]:
    if sheet not in wb.sheetnames:
        errors.append(f"Thiếu sheet '{sheet}'.")
        return []
    rows = wb[sheet].iter_rows(values_only=True)
    header = next(rows, None) or ()
    names = [str(h).strip() if h is not None else "" for h in header]
    missing = [c for c in required if c not in names]
    if missing:
        errors.append(f"Sheet {sheet}, dòng 1: thiếu cột {', '.join(missing)}.")
        return []
    known = {*required, *optional}
    result: list[_Row] = []
    for number, values in enumerate(rows, start=2):
        record = {n: _clean(v) for n, v in zip(names, values, strict=False) if n in known}
        if all(v is None for v in record.values()):
            continue
        empty = [c for c in required if record.get(c) is None]
        if empty:
            errors.append(f"Sheet {sheet}, dòng {number}: thiếu giá trị {', '.join(empty)}.")
            continue
        result.append((number, record))
    return result


def _int(value: Any, sheet: str, row: int, column: str, minimum: int, errors: list[str]) -> int:
    number: int | None = None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int) and not isinstance(value, bool):
        number = value
    elif isinstance(value, str) and value.lstrip("-").isdigit():
        number = int(value)
    if number is None or number < minimum:
        errors.append(f"Sheet {sheet}, dòng {row}: cột {column} phải là số nguyên >= {minimum}.")
        return minimum
    return number


def _check_unique(sheet: str, rows: list[_Row], errors: list[str]) -> set[str]:
    seen: set[str] = set()
    for number, record in rows:
        key = str(record["id"])
        if key in seen:
            errors.append(f"Sheet {sheet}, dòng {number}: trùng id '{key}'.")
        seen.add(key)
    return seen


class OpenpyxlExcelImporter:
    def __init__(self) -> None:
        # Cảnh báo của lần đọc gần nhất (ví dụ cột chưa được áp dụng vào mô hình).
        self.warnings: list[str] = []

    def load_data(self, path: Path) -> dict[str, Any]:
        try:
            wb = load_workbook(path, data_only=True)
        except (OSError, InvalidFileException, zipfile.BadZipFile, KeyError) as exc:
            raise ExcelImportError(f"Không đọc được file Excel {path}: {exc}") from exc

        self.warnings = []
        if is_english_workbook(wb.sheetnames) and SHEET_ROOMS not in wb.sheetnames:
            data, self.warnings = convert(wb)
            return data

        errors: list[str] = []
        rooms = _read_sheet(
            wb, SHEET_ROOMS, ("id", "name", "capacity"), ("room_type", "also_supports"), errors
        )
        teachers = _read_sheet(wb, SHEET_TEACHERS, ("id", "name"), (), errors)
        classes = _read_sheet(wb, SHEET_CLASSES, ("id", "name", "size"), ("kind",), errors)
        courses = _read_sheet(
            wb,
            SHEET_COURSES,
            ("id", "name", "class_id", "teacher_id", "periods_per_week"),
            ("required_room_type",),
            errors,
        )

        _check_unique(SHEET_ROOMS, rooms, errors)
        teacher_ids = _check_unique(SHEET_TEACHERS, teachers, errors)
        class_ids = _check_unique(SHEET_CLASSES, classes, errors)
        _check_unique(SHEET_COURSES, courses, errors)

        room_data = []
        for n, r in rooms:
            supports = [s.strip() for s in str(r.get("also_supports") or "").split(",")]
            room_data.append(
                {
                    "id": str(r["id"]),
                    "name": str(r["name"]),
                    "capacity": _int(r["capacity"], SHEET_ROOMS, n, "capacity", 1, errors),
                    "room_type": str(r.get("room_type") or "standard"),
                    "also_supports": [s for s in supports if s],
                }
            )

        class_data = [
            {
                "id": str(r["id"]),
                "name": str(r["name"]),
                "size": _int(r["size"], SHEET_CLASSES, n, "size", 1, errors),
                "kind": str(r.get("kind") or "regular"),
            }
            for n, r in classes
        ]
        course_data = []
        for n, r in courses:
            if str(r["class_id"]) not in class_ids:
                errors.append(f"Sheet {SHEET_COURSES}, dòng {n}: lớp '{r['class_id']}' không có.")
            if str(r["teacher_id"]) not in teacher_ids:
                errors.append(
                    f"Sheet {SHEET_COURSES}, dòng {n}: giáo viên '{r['teacher_id']}' không có."
                )
            periods = _int(r["periods_per_week"], SHEET_COURSES, n, "periods_per_week", 1, errors)
            course_data.append(
                {
                    "id": str(r["id"]),
                    "name": str(r["name"]),
                    "class_id": str(r["class_id"]),
                    "teacher_id": str(r["teacher_id"]),
                    "periods_per_week": periods,
                    "required_room_type": str(r.get("required_room_type") or "standard"),
                }
            )

        if errors:
            raise ExcelImportError("\n".join(errors))

        data = {
            "schema_version": SCHEMA_VERSION,
            "grid": {
                "days": NUM_DAYS,
                "periods": NUM_PERIODS,
                "allowed_periods": list(ALLOWED_PERIODS),
            },
            "rooms": room_data,
            "teachers": [{"id": str(r["id"]), "name": str(r["name"])} for _, r in teachers],
            "class_groups": class_data,
            "courses": course_data,
            "constraints": {"max_consecutive_periods": MAX_CONSECUTIVE_PERIODS},
        }
        return data

    def load(self, path: Path) -> Problem:
        data = self.load_data(path)
        try:
            return problem_from_dict(data)
        except (ValueError, TypeError) as exc:
            raise ExcelImportError(f"Dữ liệu không hợp lệ: {exc}") from exc
